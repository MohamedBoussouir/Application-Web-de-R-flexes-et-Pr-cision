"""
Génération automatique du PowerPoint (.pptx) - AXE 3
Thème : Vocabulaire à Distinguer ([[Prototype]] vs .prototype vs __proto__)
Style : Modern Dark (Deep Navy #0A0E17, Cyan #00F2FE, Purple #9B51E0, Rose #FF4B6E)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_axis3():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # الألوان المتناسقة المعتمدة
    BG_DARK     = RGBColor(10, 14, 23)     # #0A0E17
    CARD_BG     = RGBColor(20, 26, 43)     # #141A2B
    CARD_BORDER = RGBColor(38, 48, 77)     # #26304D
    CODE_BG     = RGBColor(6, 9, 15)       # #06090F
    CYAN        = RGBColor(0, 242, 254)    # #00F2FE
    PURPLE      = RGBColor(155, 81, 224)   # #9B51E0
    ROSE        = RGBColor(255, 75, 110)   # #FF4B6E
    GREEN       = RGBColor(0, 245, 160)    # #00F5A0
    TEXT_WHITE  = RGBColor(241, 245, 249)  # #F1F5F9
    TEXT_MUTED  = RGBColor(148, 163, 184)  # #94A3B8

    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()

    def add_header(slide, badge, title, subtitle):
        # Badge
        tb_badge = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(6.0), Inches(0.35))
        tf_b = tb_badge.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        p_b.text = f"●  {badge.upper()}"
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = CYAN
        p_b.font.name = "Arial"

        # Titre
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(11.7), Inches(0.65))
        tf_t = tb_title.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.font.name = "Arial"

        # Sous-titre
        if subtitle:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.48), Inches(11.7), Inches(0.45))
            tf_s = tb_sub.text_frame
            tf_s.word_wrap = True
            p_s = tf_s.paragraphs[0]
            p_sub = tf_s.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_MUTED
            p_sub.font.name = "Calibri"

    def add_card(slide, left, top, w, h, title, bullets, col):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), w - Inches(0.4), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(16)
        pt.font.bold = True
        pt.font.color.rgb = col
        pt.font.name = "Arial"
        pt.space_after = Pt(12)

        for b in bullets:
            p = tf.add_paragraph()
            p.text = f"•  {b}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_WHITE
            p.font.name = "Calibri"
            p.space_after = Pt(6)

    def add_code(slide, left, top, w, h, lines):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = CODE_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.18), w - Inches(0.5), h - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True

        first = True
        for line in lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = line
            p.font.size = Pt(11.5)
            p.font.name = "Consolas"
            if line.strip().startswith("//"):
                p.font.color.rgb = TEXT_MUTED
            elif "const" in line or "function" in line or "===" in line:
                p.font.color.rgb = CYAN
            elif "true" in line or "undefined" in line:
                p.font.color.rgb = GREEN
            else:
                p.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE AXE 3 : LE VOCABULAIRE CLÉ À DISTINGUER
    # =========================================================================
    s = prs.slides.add_slide(blank_layout)
    apply_bg(s)
    add_header(s, "Axe 3 • Rigueur Terminologique",
               "Vocabulaire à Distinguer : Les 3 Notions Clés",
               "Ne plus jamais confondre le lien interne, le moule du constructeur et l'accesseur historique.")

    # 3 بطاقات للمصطلحات الثلاثة
    add_card(s, Inches(0.8), Inches(2.2), Inches(3.7), Inches(2.8),
             "1. [[Prototype]]", [
                 "Pointeur interne secret du moteur JS.",
                 "Existe dans TOUS les objets.",
                 "Relie physiquement l'objet à son parent.",
                 "Accès officiel et standard :",
                 "Object.getPrototypeOf(obj)"
             ], CYAN)

    add_card(s, Inches(4.8), Inches(2.2), Inches(3.7), Inches(2.8),
             "2. Fonction.prototype", [
                 "Objet ordinaire en mémoire.",
                 "Existe UNIQUEMENT sur constructeurs & classes.",
                 "Sert de 'moule' pour les futures instances.",
                 "Devient le [[Prototype]] créé par 'new'.",
                 "User.prototype.sayHi = ..."
             ], PURPLE)

    add_card(s, Inches(8.8), Inches(2.2), Inches(3.7), Inches(2.8),
             "3. __proto__", [
                 "Accesseur historique (getter / setter).",
                 "Hérité depuis Object.prototype.",
                 "Exposé jadis par les navigateurs.",
                 "Officiellement DÉPRÉCIÉ aujourd'hui.",
                 "À proscrire dans du code moderne."
             ], ROSE)

    # كود الإثبات أسفل البطاقات
    code_lines = [
        "function User(name) { this.name = name; }",
        "User.prototype.sayHi = function() {}; // 1. .prototype existe uniquement sur le constructeur",
        "const u1 = new User(\"Mohamed\");",
        "",
        "// 2. [[Prototype]] de l'instance pointe vers User.prototype :",
        "console.log(Object.getPrototypeOf(u1) === User.prototype); // true",
        "console.log(u1.prototype);                                // undefined (l'instance n'en a pas !)",
        "",
        "// 3. __proto__ est l'ancien chemin déprécié pointant vers le même lien :",
        "console.log(u1.__proto__ === User.prototype);              // true (À éviter en production)"
    ]
    add_code(s, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.8), code_lines)

    # سيناريو الإلقاء في خانة الملاحظات
    s.notes_slide.notes_text_frame.text = (
        "Script oral pour l'Axe 3 :\n"
        "Pour bien maîtriser les prototypes, il faut absolument distinguer 3 notions souvent confondues :\n"
        "1. [[Prototype]] : c'est le lien interne et invisible présent dans TOUS les objets, qui pointe vers le parent. On le lit avec Object.getPrototypeOf(obj).\n"
        "2. .prototype : c'est une propriété réservée aux constructeurs et classes. C'est le moule des futures instances créées avec 'new'.\n"
        "3. __proto__ : c'est un accesseur historique créé par les navigateurs. Il est aujourd'hui totalement déprécié."
    )

    output_filename = "Axe3_Vocabulaire_Presentation.pptx"
    prs.save(output_filename)
    print(f"[OK] Présentation générée avec succès : {output_filename}")

if __name__ == "__main__":
    build_axis3()